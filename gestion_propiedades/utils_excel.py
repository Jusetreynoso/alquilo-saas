import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import io

def generar_excel_movimientos_detallados(movimientos, total_ingresos, total_gastos, balance_neto, fecha_inicio, fecha_fin):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Movimientos Detallados"
    
    # Mostrar líneas de cuadrícula
    ws.views.sheetView[0].showGridLines = True
    
    # Encabezado Principal del Documento
    ws.merge_cells('A1:G1')
    title_cell = ws['A1']
    title_cell.value = "REPORTE DETALLADO DE INGRESOS Y GASTOS"
    title_cell.font = Font(name='Calibri', size=15, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color='1A252F', end_color='1A252F', fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32

    # Subtítulo Rango de Fechas
    ws.merge_cells('A2:G2')
    sub_cell = ws['A2']
    sub_cell.value = f"Período del {fecha_inicio.strftime('%d/%m/%Y')} al {fecha_fin.strftime('%d/%m/%Y')}"
    sub_cell.font = Font(name='Calibri', size=11, italic=True, color='333333')
    sub_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 20

    ws.append([]) # Fila 3 en blanco

    # Encabezados de Tabla
    headers = ["Fecha", "Propiedad", "Propietario", "Inquilino / Referencia", "Tipo", "Concepto / Descripción", "Monto (RD$)"]
    ws.append(headers)
    
    header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws.row_dimensions[4].height = 25

    # Filas de Datos
    ingreso_font = Font(color='27AE60', bold=True)
    gasto_font = Font(color='C0392B', bold=True)
    
    for row_idx, m in enumerate(movimientos, start=5):
        fecha_str = m['fecha'].strftime('%d/%m/%Y') if hasattr(m['fecha'], 'strftime') else str(m['fecha'])
        prop_str = f"{m['propiedad']} ({m['grupo_o_residencial']})" if m.get('grupo_o_residencial') else m['propiedad']
        
        ws.append([
            fecha_str,
            prop_str,
            m['propietario'],
            m['persona'],
            m['tipo'],
            m['concepto'],
            float(m['monto'])
        ])
        
        row_cells = ws[row_idx]
        for c_idx, cell in enumerate(row_cells, start=1):
            cell.border = thin_border
            if c_idx == 1:
                cell.alignment = Alignment(horizontal='center')
            elif c_idx == 5:
                cell.alignment = Alignment(horizontal='center')
                if m['tipo'] == 'INGRESO':
                    cell.font = Font(color='27AE60', bold=True)
                else:
                    cell.font = Font(color='C0392B', bold=True)
            elif c_idx == 7:
                cell.number_format = '$#,##0.00;($#,##0.00);"-"'
                cell.alignment = Alignment(horizontal='right')
                if m['tipo'] == 'INGRESO':
                    cell.font = ingreso_font
                else:
                    cell.font = gasto_font

    # Fila de Totales
    current_row = len(movimientos) + 6
    
    ws.cell(row=current_row, column=6, value="TOTAL INGRESOS:").font = Font(bold=True)
    ws.cell(row=current_row, column=6).alignment = Alignment(horizontal='right')
    tot_ing = ws.cell(row=current_row, column=7, value=float(total_ingresos))
    tot_ing.font = Font(color='27AE60', bold=True)
    tot_ing.number_format = '$#,##0.00'
    
    current_row += 1
    ws.cell(row=current_row, column=6, value="TOTAL GASTOS:").font = Font(bold=True)
    ws.cell(row=current_row, column=6).alignment = Alignment(horizontal='right')
    tot_gas = ws.cell(row=current_row, column=7, value=float(total_gastos))
    tot_gas.font = Font(color='C0392B', bold=True)
    tot_gas.number_format = '$#,##0.00'

    current_row += 1
    ws.cell(row=current_row, column=6, value="BALANCE NETO:").font = Font(bold=True, size=12)
    ws.cell(row=current_row, column=6).alignment = Alignment(horizontal='right')
    tot_net = ws.cell(row=current_row, column=7, value=float(balance_neto))
    tot_net.font = Font(color='1A252F', bold=True, size=12)
    tot_net.number_format = '$#,##0.00'

    # Autoajustar ancho de columnas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if cell.row > 2 and len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def generar_excel_estado_propiedades_propietario(data):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estado Propiedades"
    
    ws.views.sheetView[0].showGridLines = True
    
    # Encabezado Principal del Documento
    ws.merge_cells('A1:J1')
    title_cell = ws['A1']
    title_cell.value = f"REPORTE MENSUAL DE ESTADO DE PROPIEDADES ({data['nombre_mes'].upper()} {data['anio']})"
    title_cell.font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color='1A252F', end_color='1A252F', fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    ws.append([]) # Fila 2 vacía

    # Encabezados de Tabla
    headers = [
        "Propiedad", "Propietario", "Inquilino", "Renta Pactada (RD$)",
        "Estado Pago", "Fecha Pago", "Días Atraso", "Monto Cobrado (RD$)",
        "Gastos Directos (RD$)", "Neto Propietario (RD$)"
    ]
    ws.append(headers)
    
    header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws.row_dimensions[3].height = 24

    for row_idx, item in enumerate(data['reporte_propiedades'], start=4):
        prop_str = f"{item['propiedad_nombre']} ({item['grupo_o_residencial']})" if item.get('grupo_o_residencial') else item['propiedad_nombre']
        ws.append([
            prop_str,
            item['propietario'],
            item['inquilino'],
            float(item['renta_pactada']),
            item['estado_label'],
            item['fecha_pago'] or '-',
            item['dias_retraso'] if item['dias_retraso'] > 0 else '-',
            float(item['monto_pagado']),
            float(item['total_gastos']),
            float(item['monto_neto'])
        ])
        
        row_cells = ws[row_idx]
        for c_idx, cell in enumerate(row_cells, start=1):
            cell.border = thin_border
            if c_idx in [4, 8, 9, 10]:
                cell.number_format = '$#,##0.00;($#,##0.00);"-"'
                cell.alignment = Alignment(horizontal='right')
            elif c_idx in [5, 6, 7]:
                cell.alignment = Alignment(horizontal='center')

    # Fila Totales
    last_row = len(data['reporte_propiedades']) + 5
    tot = data['totales_resumen']
    
    ws.cell(row=last_row, column=3, value="TOTALES GENERALES:").font = Font(bold=True)
    ws.cell(row=last_row, column=3).alignment = Alignment(horizontal='right')
    
    c_pact = ws.cell(row=last_row, column=4, value=float(tot['total_renta_pactada']))
    c_pact.font = Font(bold=True)
    c_pact.number_format = '$#,##0.00'

    c_cob = ws.cell(row=last_row, column=8, value=float(tot['total_cobrado']))
    c_cob.font = Font(bold=True, color='27AE60')
    c_cob.number_format = '$#,##0.00'

    c_gas = ws.cell(row=last_row, column=9, value=float(tot['total_gastos']))
    c_gas.font = Font(bold=True, color='C0392B')
    c_gas.number_format = '$#,##0.00'

    c_net = ws.cell(row=last_row, column=10, value=float(tot['neto_propietario']))
    c_net.font = Font(bold=True, size=11, color='1A252F')
    c_net.number_format = '$#,##0.00'

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
